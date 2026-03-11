#!/usr/bin/env python3
"""
Wakeup Log Parser
-----------------
Traverses ./logs/<device>/hilog/ directories, extracts wakeup event timestamps
from Android (hiapplogcat-log) and HarmonyOS (hilog.*.txt) log files,
and writes the results into an Excel workbook with one sheet per device.

HarmonyOS log format (hilog.*.txt):
  T1  ===de / ===detected   (Level-1 wakeup entry)
  T2  ===start::            (Level-2 wakeup entry)
  T3  onResult:: isShouldResponse  (Decision log)

Android log format (hiapplogcat-log*):
  No T1 (level-1 wakeup does not exist on Android)
  T2  ===start::            (Level-2 wakeup entry)
  T3  onResult:: isShouldResponse  (Decision log)
"""

import os
import re
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ── patterns ──────────────────────────────────────────────────────────────────

# Timestamp at the start of a log line: MM-DD HH:MM:SS.mmm
TS_RE = re.compile(r"(\d{2}-\d{2} \d{2}:\d{2}:\d{2}\.\d{3})")

# Keyword patterns
RE_T1 = re.compile(r"===de")                         # level-1 wakeup (detected)
RE_T2 = re.compile(r"===start::")                    # level-2 wakeup (start)
RE_T3 = re.compile(r"onResult::\s*isShouldResponse") # decision log

# Extract the boolean value after "isShouldResponse ="
RE_DECISION = re.compile(r"isShouldResponse\s*=\s*(true|false)", re.IGNORECASE)

# ── helpers ───────────────────────────────────────────────────────────────────

def extract_timestamp(line: str) -> str | None:
    """Return the first MM-DD HH:MM:SS.mmm timestamp found in *line*, or None."""
    m = TS_RE.search(line)
    return m.group(1) if m else None


def parse_log_file(filepath: Path) -> list[dict]:
    """
    Parse a single log file and return a list of wakeup event groups.
    Each group is a dict with keys: t1, t2, t3, decision.

    Grouping rules:
      HarmonyOS path (hilog.*.txt) — has T1:
        T1 → opens a new group (flushes any incomplete previous group)
        T2 → attaches to the open group
        T3 → closes and saves the group

      Android path (hiapplogcat-log*) — no T1:
        T2 → opens a new group (because there will never be a T1)
             if a group is already open, flush it first
        T3 → closes and saves the group
    """
    groups: list[dict] = []
    pending: dict | None = None  # current open group

    try:
        with open(filepath, "r", encoding="utf-8", errors="replace") as fh:
            for line in fh:
                if RE_T1.search(line):
                    ts = extract_timestamp(line)
                    if ts:
                        if pending:
                            groups.append(pending)
                        pending = {"t1": ts, "t2": None, "t3": None, "decision": None}

                elif RE_T2.search(line):
                    ts = extract_timestamp(line)
                    if not ts:
                        continue
                    if pending is None:
                        # Android: no T1 ever appears → T2 opens the group
                        pending = {"t1": None, "t2": ts, "t3": None, "decision": None}
                    elif pending["t2"] is None:
                        # HarmonyOS: T2 follows T1 in the same group
                        pending["t2"] = ts
                    else:
                        # A second T2 arrived before T3 — flush the stale group
                        # and start fresh (handles back-to-back wakeups on Android)
                        groups.append(pending)
                        pending = {"t1": None, "t2": ts, "t3": None, "decision": None}

                elif RE_T3.search(line):
                    ts = extract_timestamp(line)
                    if not ts:
                        continue
                    if pending is None:
                        # T3 without any preceding T1/T2 — create a minimal group
                        pending = {"t1": None, "t2": None, "t3": None, "decision": None}
                    if pending["t3"] is None:
                        pending["t3"] = ts
                        m = RE_DECISION.search(line)
                        pending["decision"] = m.group(1).lower() if m else "unknown"
                        groups.append(pending)
                        pending = None

    except OSError as exc:
        print(f"  [WARN] Cannot read {filepath}: {exc}", file=sys.stderr)

    # Flush any trailing incomplete group (e.g. log cut off before T3)
    if pending:
        groups.append(pending)

    return groups


def is_log_file(path: Path) -> bool:
    """Return True if *path* looks like an Android or HarmonyOS log file."""
    name = path.name
    # Android: hiapplogcat-log (no extension, or with suffix like hiapplogcat-log.1)
    if name.startswith("hiapplogcat-log"):
        return True
    # HarmonyOS: hilog.xxxx.txt
    if name.startswith("hilog") and name.endswith(".txt"):
        return True
    return False


# ── Excel helpers ─────────────────────────────────────────────────────────────

HEADER = ["一级唤醒时间 (T1)", "二级唤醒时间 (T2)", "决策时间 (T3)", "决策结果（是否响应）"]

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
CELL_FONT   = Font(name="Calibri", size=10)
CENTER      = Alignment(horizontal="center", vertical="center")


def style_sheet(ws):
    """Apply header styles and column widths to a worksheet."""
    for col_idx, title in enumerate(HEADER, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = CENTER

    # Fixed column widths
    col_widths = [22, 22, 22, 22]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 20


def write_events_to_sheet(ws, events: list[dict]):
    """Write event rows starting at row 2."""
    for row_idx, ev in enumerate(events, start=2):
        values = [ev["t1"], ev["t2"], ev["t3"], ev["decision"]]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val if val else "—")
            cell.font      = CELL_FONT
            cell.alignment = CENTER


# ── main ──────────────────────────────────────────────────────────────────────

def main():
    logs_root = Path("./logs")
    if not logs_root.is_dir():
        print(f"[ERROR] Directory not found: {logs_root.resolve()}", file=sys.stderr)
        sys.exit(1)

    wb = Workbook()
    # Remove the default empty sheet
    wb.remove(wb.active)

    device_dirs = sorted(
        d for d in logs_root.iterdir()
        if d.is_dir()
    )

    if not device_dirs:
        print("[WARN] No device directories found under ./logs/")
        sys.exit(0)

    total_events = 0

    for device_dir in device_dirs:
        device_name = device_dir.name
        hilog_dir   = device_dir / "hilog"

        if not hilog_dir.is_dir():
            print(f"[SKIP] {device_name}: no hilog/ subdirectory")
            continue

        log_files = sorted(f for f in hilog_dir.iterdir() if f.is_file() and is_log_file(f))

        if not log_files:
            print(f"[SKIP] {device_name}: no recognised log files in hilog/")
            continue

        print(f"[INFO] Processing device: {device_name} ({len(log_files)} file(s))")

        all_events: list[dict] = []
        for lf in log_files:
            events = parse_log_file(lf)
            print(f"  {lf.name}: {len(events)} wakeup event(s)")
            all_events.extend(events)

        # Excel sheet name: max 31 chars, no special chars
        sheet_name = re.sub(r"[\\/*?\[\]:]", "_", device_name)[:31]
        ws = wb.create_sheet(title=sheet_name)
        style_sheet(ws)
        write_events_to_sheet(ws, all_events)
        total_events += len(all_events)
        print(f"  → Sheet '{sheet_name}': {len(all_events)} row(s) written")

    if not wb.sheetnames:
        print("[WARN] No data extracted – Excel file will not be created.")
        sys.exit(0)

    output_path = Path("wakeup_events.xlsx")
    wb.save(output_path)
    print(f"\n[DONE] Saved {total_events} total event(s) to: {output_path.resolve()}")


if __name__ == "__main__":
    main()
